from pathlib import Path

from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import formatters
from .fields import insert_field, insert_page, insert_seq


class DocumentBuilder:
    """University-agnostic DOCX primitives. Rules come from config."""

    def __init__(self, document, config: dict):
        self.doc = document
        self.config = config
        self.current_chapter = 0

    def prepare(self):
        formatters.setup_true_multilevel_list(self.doc)
        formatters.cleanup_template(self.doc)
        formatters.configure_body_style(self.doc, self.config)
        self._setup_title_cover_section()
        formatters.configure_heading_styles(self.doc, self.config)
        formatters.configure_front_matter_heading_style(self.doc, self.config)
        formatters.configure_toc_styles(self.doc)

    def finalize(self):
        formatters.configure_body_style(self.doc, self.config)
        formatters.configure_heading_styles(self.doc, self.config)
        formatters.configure_front_matter_heading_style(self.doc, self.config)
        formatters.configure_toc_styles(self.doc)
        formatters.enforce_body_paragraphs(self.doc, self.config)

    def _heading_size(self, level):
        headings = self.config.get("styles", {}).get("headings", {})
        cfg = headings.get(f"JUW_H{level}", {})
        return float(cfg.get("size", 20 - ((level - 1) * 2)))

    def force_times_new_roman(self, run):
        run.font.name = "Times New Roman"
        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        r_fonts.set(qn("w:ascii"), "Times New Roman")
        r_fonts.set(qn("w:hAnsi"), "Times New Roman")
        r_fonts.set(qn("w:eastAsia"), "Times New Roman")
        r_fonts.set(qn("w:cs"), "Times New Roman")

    def add_bottom_border(self, paragraph):
        paragraph.paragraph_format.space_after = Pt(12)
        p_pr = paragraph._element.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "24")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "000000")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    def _exclude_from_toc(self, paragraph):
        """Keep Heading-like visuals, but outline level Body Text so TOC \\o skips it."""
        p_pr = paragraph._element.get_or_add_pPr()
        existing = p_pr.find(qn("w:outlineLvl"))
        if existing is not None:
            p_pr.remove(existing)
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "9")
        p_pr.append(outline)

    def _apply_front_matter_heading(self, paragraph):
        paragraph.style = self.doc.styles["JUW Front Matter"]
        self._exclude_from_toc(paragraph)

    def add_heading(self, text, level):
        paragraph = self.doc.add_paragraph()
        paragraph.style = self.doc.styles[f"Heading {level}"]
        paragraph_format = paragraph.paragraph_format

        if level == 1:
            paragraph_format.space_before = Pt(0)
            paragraph_format.space_after = Pt(0)
        elif level == 2:
            paragraph_format.space_before = Pt(18)
            paragraph_format.space_after = Pt(6)

        if self.current_chapter > 0 and level > 1:
            p_pr = paragraph._element.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            ilvl_el = OxmlElement("w:ilvl")
            ilvl_el.set(qn("w:val"), str(level - 1))
            num_id_el = OxmlElement("w:numId")
            num_id_el.set(qn("w:val"), "999")
            num_pr.append(ilvl_el)
            num_pr.append(num_id_el)
            p_pr.append(num_pr)

        run = paragraph.add_run(text)
        self.force_times_new_roman(run)
        size_pt, all_caps, small_caps = formatters.juw_heading_appearance(text, level)
        formatters.apply_heading_run(run, size_pt, bold=True, all_caps=all_caps, small_caps=small_caps)

        if level == 1 and text.upper().startswith("CHAPTER"):
            self.add_bottom_border(paragraph)
        return paragraph

    def add_body_text(self, text, level=2):
        if not text or not str(text).strip():
            return None

        paragraph = self.doc.add_paragraph()
        paragraph.style = self.doc.styles["Normal"]
        run = paragraph.add_run(str(text).strip())
        body = self.config.get("styles", {}).get("body", {})
        formatters.apply_body_run(run, body.get("size", 12))
        run.font.color.rgb = RGBColor(0, 0, 0)

        formatters.apply_body_paragraph(paragraph, self.config, document=self.doc)
        formatters.apply_juw_body_indent(paragraph, level)
        return paragraph

    def add_paragraphs(self, text, level=2):
        for block in split_paragraphs(text):
            self.add_body_text(block, level=level)

    def add_bullet_point(self, text, level=1, heading_level=2):
        """Foundational JUW bullet: glyph + hanging indent, TNR 12, 1.5, justified."""
        if not text or not str(text).strip():
            return None
        bullet_chars = {1: "\u2022", 2: "o", 3: "\u25AA"}
        return self._add_list_item(
            str(text).strip(),
            marker=f"{bullet_chars.get(level, chr(8226))}\t",
            level=level,
            heading_level=heading_level,
        )

    def add_numbered_item(self, text, number, level=1, heading_level=2):
        if not text or not str(text).strip():
            return None
        label = str(number).strip()
        marker = f"{label}.\t" if "." not in label else f"{label}\t"
        if label.endswith("."):
            marker = f"{label}\t"
        return self._add_list_item(
            str(text).strip(),
            marker=marker,
            level=level,
            heading_level=heading_level,
        )

    def add_bullet_list(self, items, level=1, heading_level=2):
        return self.add_body_list(items, kind="bullet", default_level=level, heading_level=heading_level)

    def add_numbered_list(self, items, level=1, heading_level=2):
        return self.add_body_list(items, kind="number", default_level=level, heading_level=heading_level)

    def add_body_list(self, items, kind="bullet", default_level=1, heading_level=2):
        """Render a structured list. Numbered items use 1 / 1.1 / 1.2 / 2, not heading numId 999."""
        normalized = []
        previous_level = 1
        for raw in items or []:
            if isinstance(raw, dict):
                text = str(raw.get("text") or raw.get("title") or "").strip()
                level = int(raw.get("level") or default_level)
            elif isinstance(raw, (tuple, list)) and raw:
                text = str(raw[0]).strip()
                level = int(raw[1]) if len(raw) > 1 else default_level
            else:
                text = str(raw).strip()
                level = default_level
            if not text:
                continue
            level = max(1, min(int(level), 4))
            if not normalized:
                level = 1
            else:
                level = min(level, previous_level + 1)
            normalized.append((text, level))
            previous_level = level

        numbered = str(kind or "bullet").lower() in {"number", "numbered", "ol", "ordered"}
        counters = [0, 0, 0, 0]
        paragraphs = []
        for text, level in normalized:
            counters[level - 1] += 1
            for index in range(level, 4):
                counters[index] = 0
            if numbered:
                label = ".".join(str(counters[index]) for index in range(level))
                marker = f"{label}.\t" if level == 1 else f"{label}\t"
                paragraph = self._add_list_item(text, marker, level=level, heading_level=heading_level)
            else:
                paragraph = self.add_bullet_point(text, level=level, heading_level=heading_level)
            if paragraph is not None:
                paragraphs.append(paragraph)
        return paragraphs

    def _add_list_item(self, text, marker, level=1, heading_level=2):
        paragraph = self.doc.add_paragraph()
        paragraph.style = self.doc.styles["Normal"]
        base_indent = 0.5 if heading_level <= 2 else 0.75
        paragraph_format = paragraph.paragraph_format
        paragraph_format.left_indent = Inches(base_indent + (0.25 * (max(level, 1) - 1)))
        paragraph_format.first_line_indent = Inches(-0.20)
        paragraph_format.line_spacing = self.config.get("layout", {}).get("spacing", {}).get("line_spacing", 1.5)
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        try:
            paragraph_format.tab_stops.add_tab_stop(paragraph_format.left_indent)
        except Exception:
            pass

        p_pr = paragraph._element.get_or_add_pPr()
        existing = p_pr.find(qn("w:spacing"))
        if existing is not None:
            p_pr.remove(existing)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:beforeAutospacing"), "0")
        spacing.set(qn("w:afterAutospacing"), "0")
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        snap = p_pr.find(qn("w:snapToGrid"))
        if snap is None:
            snap = OxmlElement("w:snapToGrid")
            p_pr.append(snap)
        snap.set(qn("w:val"), "0")
        jc = p_pr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            p_pr.append(jc)
        jc.set(qn("w:val"), "both")

        body = self.config.get("styles", {}).get("body", {})
        size_pt = body.get("size", 12)
        marker_run = paragraph.add_run(marker)
        formatters.apply_body_run(marker_run, size_pt)
        marker_run.font.color.rgb = RGBColor(0, 0, 0)
        text_run = paragraph.add_run(text)
        formatters.apply_body_run(text_run, size_pt)
        text_run.font.color.rgb = RGBColor(0, 0, 0)
        return paragraph

    def add_front_matter_page(self, title, text="", page_break_before=False):
        previous_chapter = self.current_chapter
        self.current_chapter = 0

        paragraph = self.doc.add_paragraph()
        self._apply_front_matter_heading(paragraph)
        if page_break_before:
            paragraph.paragraph_format.page_break_before = True
        self.add_bottom_border(paragraph)
        run = paragraph.add_run(title.upper())
        self.force_times_new_roman(run)
        run.font.size = Pt(self._heading_size(1))
        run.bold = True

        self.add_paragraphs(text, level=1)
        self.doc.add_page_break()
        self.current_chapter = previous_chapter

    def add_title_page(self, meta=None):
        """
        Main cover page matching FYP_Report_Complete_Format-2026 (page 1).
        """
        meta = meta or {}
        previous_chapter = self.current_chapter
        self.current_chapter = 0
        title_cfg = self.config.get("layout", {}).get("title_page", {})

        project_title = str(
            meta.get("projectTitle") or meta.get("title") or "PROJECT TITLE"
        ).strip()
        submission = str(
            meta.get("submissionMonthYear") or meta.get("submissionDate") or "Month 20XX"
        ).strip()
        report_label = str(
            meta.get("reportLabel")
            or title_cfg.get("report_label")
            or "BS (CS) Final Year Project Report"
        ).strip()
        department_line = str(
            title_cfg.get("department_line")
            or "Department of Computer Science and Software Engineering"
        )
        university_line = str(
            title_cfg.get("university_line") or "Jinnah University for Women"
        )
        address_line = str(
            title_cfg.get("address_line") or "5-C Nazimabad, Karachi 74600"
        )
        student_tabs = title_cfg.get("student_tabs_inches") or [2.0, 4.5]

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = 1.0

        title_size = float(title_cfg.get("title_font_size", 26))
        title_lines = self._split_title_lines(project_title, max_lines=2)
        for index, line in enumerate(title_lines):
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(
                6 if index < len(title_lines) - 1 else 46
            )
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(line)
            self._apply_run_style(run, size=title_size, bold=False, small_caps=True)

        logo_w = float(title_cfg.get("logo_width_inches", 1.24))
        logo_h = float(title_cfg.get("logo_height_inches", 1.34))
        self._add_logo_or_reserved_space(
            meta,
            width_inches=logo_w,
            height_inches=logo_h,
            space_before_pt=6,
            space_after_pt=22,
        )

        self._add_centered_line(report_label, size=16, bold=True, space_after=28, line_spacing=1.0)

        submitted = self.doc.add_paragraph()
        submitted.alignment = WD_ALIGN_PARAGRAPH.LEFT
        submitted.paragraph_format.space_before = Pt(6)
        submitted.paragraph_format.space_after = Pt(6)
        submitted.paragraph_format.line_spacing = 1.0
        run = submitted.add_run("Submitted by")
        self._apply_run_style(run, size=13, bold=True)

        self._add_student_tab_rows(
            self._normalize_students(meta),
            tab_positions_inches=student_tabs,
            font_size=13,
            header_bold=True,
            placeholder=True,
            seat_header="Seat number",
        )

        date_spacer = self.doc.add_paragraph()
        date_spacer.paragraph_format.space_before = Pt(36)
        date_spacer.paragraph_format.space_after = Pt(0)
        self._add_centered_line(submission, size=11, bold=True, space_after=48, line_spacing=1.0)

        self._add_centered_line(
            department_line, size=14, bold=True, space_after=2, line_spacing=1.0, small_caps=True
        )
        self._add_centered_line(university_line, size=18, bold=True, space_after=2, line_spacing=1.0)
        self._add_centered_line(
            address_line, size=12, bold=True, space_after=0, line_spacing=1.0, small_caps=True
        )

        self.current_chapter = previous_chapter

    def add_second_cover_page(self, meta=None):
        """
        Second cover / Project Advisor page matching FYP_Report_Complete_Format-2026 (page 2).
        No university logo on this page.
        """
        meta = meta or {}
        previous_chapter = self.current_chapter
        self.current_chapter = 0
        title_cfg = self.config.get("layout", {}).get("title_page", {})
        second_cfg = self.config.get("layout", {}).get("second_cover", {})

        project_title = str(
            meta.get("projectTitle") or meta.get("title") or "PROJECT TITLE"
        ).strip()
        submission = str(
            meta.get("submissionMonthYear") or meta.get("submissionDate") or "Month 20XX"
        ).strip()
        report_label = str(
            meta.get("reportLabel")
            or title_cfg.get("report_label")
            or "BS (CS) Final Year Project Report"
        ).strip()
        department_line = str(
            title_cfg.get("department_line")
            or "Department of Computer Science and Software Engineering"
        )
        university_line = str(
            title_cfg.get("university_line") or "Jinnah University for Women"
        )
        address_line = str(
            title_cfg.get("address_line") or "5-C Nazimabad, Karachi 74600"
        )
        advisor = str(
            meta.get("supervisor") or meta.get("projectAdvisor") or ""
        ).strip()
        student_tabs = second_cfg.get("student_tabs_inches") or [1.5, 3.5]
        top_spacer_pt = float(second_cfg.get("top_spacer_pt", 34))
        after_title_pt = float(second_cfg.get("after_title_pt", 60))
        after_report_label_pt = float(second_cfg.get("after_report_label_pt", 54))
        before_advisor_pt = float(second_cfg.get("before_advisor_pt", 30))
        before_date_pt = float(second_cfg.get("before_date_pt", 90))
        after_date_pt = float(second_cfg.get("after_date_pt", 90))

        title_size = float(title_cfg.get("title_font_size", 26))
        title_lines = self._split_title_lines(project_title, max_lines=2)

        # Page break on an empty paragraph so top spacing is not collapsed by Word.
        break_para = self.doc.add_paragraph()
        break_para.paragraph_format.page_break_before = True
        break_para.paragraph_format.space_before = Pt(0)
        break_para.paragraph_format.space_after = Pt(0)
        break_para.paragraph_format.line_spacing = 1.0

        for index, line in enumerate(title_lines):
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(top_spacer_pt if index == 0 else 0)
            paragraph.paragraph_format.space_after = Pt(
                6 if index < len(title_lines) - 1 else after_title_pt
            )
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(line)
            self._apply_run_style(run, size=title_size, bold=False, small_caps=True)

        self._add_centered_line(
            report_label, size=16, bold=True, space_after=after_report_label_pt, line_spacing=1.0
        )

        submitted = self.doc.add_paragraph()
        submitted.alignment = WD_ALIGN_PARAGRAPH.LEFT
        submitted.paragraph_format.space_before = Pt(6)
        submitted.paragraph_format.space_after = Pt(6)
        submitted.paragraph_format.line_spacing = 1.0
        submitted.paragraph_format.left_indent = Inches(0)
        submitted.paragraph_format.first_line_indent = Inches(0)
        run = submitted.add_run("Submitted by")
        self._apply_run_style(run, size=13, bold=True)

        # Same deterministic tab stops as the official sample (Name / Enrolment / Seat).
        self._add_student_tab_rows(
            self._normalize_students(meta),
            tab_positions_inches=student_tabs,
            font_size=13,
            header_bold=True,
            placeholder=True,
            seat_header="Seat number",
        )

        # CRITICAL: Project Advisor: Name on the SAME line.
        advisor_para = self.doc.add_paragraph()
        advisor_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        advisor_para.paragraph_format.space_before = Pt(before_advisor_pt)
        advisor_para.paragraph_format.space_after = Pt(0)
        advisor_para.paragraph_format.line_spacing = 1.0
        advisor_para.paragraph_format.left_indent = Inches(0)
        advisor_para.paragraph_format.first_line_indent = Inches(0)
        label_run = advisor_para.add_run("Project Advisor: ")
        self._apply_run_style(label_run, size=13, bold=True)
        name_run = advisor_para.add_run(advisor)
        self._apply_run_style(name_run, size=13, bold=True)

        date_spacer = self.doc.add_paragraph()
        date_spacer.paragraph_format.space_before = Pt(before_date_pt)
        date_spacer.paragraph_format.space_after = Pt(0)
        self._add_centered_line(
            submission, size=11, bold=True, space_after=after_date_pt, line_spacing=1.0
        )

        self._add_centered_line(
            department_line, size=14, bold=True, space_after=2, line_spacing=1.0, small_caps=True
        )
        self._add_centered_line(university_line, size=18, bold=True, space_after=2, line_spacing=1.0)
        self._add_centered_line(
            address_line, size=12, bold=True, space_after=0, line_spacing=1.0, small_caps=True
        )

        self._begin_roman_front_matter_section()
        self.current_chapter = previous_chapter

    def add_approval_form(self, meta=None):
        """
        Project Approval page matching FYP_Report_Complete_Format-2026 (page 3).
        """
        meta = meta or {}
        previous_chapter = self.current_chapter
        self.current_chapter = 0
        approval_cfg = self.config.get("layout", {}).get("approval_page", {})
        title_cfg = self.config.get("layout", {}).get("title_page", {})

        project_title = str(
            meta.get("projectTitle") or meta.get("title") or "Title of the project"
        ).strip()
        department_line = str(
            approval_cfg.get("department_line")
            or title_cfg.get("department_line")
            or "Department of Computer Science and Software Engineering"
        )
        university_line = str(
            approval_cfg.get("university_line") or "JINNAH UNIVERSITY FOR WOMEN"
        )
        student_tabs = approval_cfg.get("student_tabs_inches") or [1.44, 3.5]

        logo_w = float(approval_cfg.get("logo_width_inches", 0.89))
        logo_h = float(approval_cfg.get("logo_height_inches", 0.96))
        # Sample logo starts at ~0.84" while left margin is 1.5" → indent table left.
        left_margin = float(
            self.config.get("layout", {}).get("margins_inches", {}).get("left", 1.5)
        )
        right_margin = float(
            self.config.get("layout", {}).get("margins_inches", {}).get("right", 1.0)
        )
        page_width = 8.5 if str(self.config.get("layout", {}).get("page_size", "Letter")).lower() == "letter" else 8.27
        logo_left = float(approval_cfg.get("logo_left_inches", 0.84))
        table_indent = logo_left - left_margin
        # Logo sits in the left margin; text cell must reach the right margin so the
        # department Small Caps line stays on one line (~5.58" in the official sample).
        text_col = max(5.58, page_width - right_margin - logo_left - logo_w)

        header = self.doc.add_table(rows=1, cols=2)
        self._set_table_borders_none(header)
        self._set_table_cell_margins(header, inches=0.0)
        self._set_table_column_widths(header, [logo_w, text_col])
        self._set_table_left_indent(header, table_indent)
        self._set_table_alignment(header, "left")

        cell_logo = header.rows[0].cells[0]
        cell_text = header.rows[0].cells[1]
        # Prevent Word from wrapping the department cell.
        tc_pr = cell_text._tc.get_or_add_tcPr()
        if tc_pr.find(qn("w:noWrap")) is None:
            tc_pr.append(OxmlElement("w:noWrap"))

        logo_path = self._resolve_logo_path(meta)
        logo_para = cell_logo.paragraphs[0]
        logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        logo_para.paragraph_format.space_after = Pt(0)
        if logo_path is not None:
            run = logo_para.add_run()
            run.add_picture(str(logo_path), width=Inches(logo_w), height=Inches(logo_h))
        else:
            logo_para.paragraph_format.space_after = Pt(logo_h * 72)

        p1 = cell_text.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.0
        r1 = p1.add_run(department_line)
        self._apply_run_style(r1, size=14, bold=True, small_caps=True)

        p2 = cell_text.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.0
        r2 = p2.add_run(university_line)
        self._apply_run_style(r2, size=16, bold=True)

        # PROJECT APPROVAL — centered under header, bold, underlined
        title_para = self.doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_para.paragraph_format.space_before = Pt(4)
        title_para.paragraph_format.space_after = Pt(28)
        title_para.paragraph_format.line_spacing = 1.0
        title_run = title_para.add_run("PROJECT APPROVAL")
        self._apply_run_style(title_run, size=14, bold=True, underline=True)

        # Project Title:    value (same line; tab to sample position)
        pt = self.doc.add_paragraph()
        pt.paragraph_format.space_after = Pt(18)
        pt.paragraph_format.line_spacing = 1.0
        pt.paragraph_format.left_indent = Inches(0)
        pt.paragraph_format.first_line_indent = Inches(0)
        self._set_paragraph_tab_stops(pt, [1.44])
        pt_run = pt.add_run(f"Project Title:\t{project_title}")
        self._apply_run_style(pt_run, size=12, bold=False)

        by_para = self.doc.add_paragraph()
        by_para.paragraph_format.space_before = Pt(6)
        by_para.paragraph_format.space_after = Pt(6)
        by_para.paragraph_format.left_indent = Inches(0)
        by_para.paragraph_format.first_line_indent = Inches(0)
        by_run = by_para.add_run("By")
        self._apply_run_style(by_run, size=12, bold=False)

        self._add_student_tab_rows(
            self._normalize_students(meta),
            tab_positions_inches=student_tabs,
            font_size=12,
            header_bold=False,
            placeholder=True,
            seat_header="Seat",
        )

        ac = self.doc.add_paragraph()
        ac.paragraph_format.space_before = Pt(48)
        ac.paragraph_format.space_after = Pt(36)
        ac_run = ac.add_run("Approval Committee:")
        self._apply_run_style(ac_run, size=12, bold=False)

        internal_name = str(
            meta.get("internalExaminer")
            or meta.get("supervisor")
            or meta.get("projectAdvisor")
            or ""
        ).strip()
        external_name = str(meta.get("externalExaminer") or "").strip()
        internal_desig = str(
            meta.get("internalExaminerDesignation")
            or meta.get("internalDesignation")
            or ""
        ).strip()
        external_desig = str(
            meta.get("externalExaminerDesignation")
            or meta.get("externalDesignation")
            or ""
        ).strip()
        external_org = str(
            meta.get("externalExaminerOrganization")
            or meta.get("externalOrganization")
            or ""
        ).strip()
        hod_name = str(meta.get("headOfDepartment") or "").strip()

        # Two-column advisor block at sample x positions (~108 and ~360).
        # Content width is 6.0" (Letter - 1.5" left - 1.0" right); cols 3.5"+2.5".
        sig = self.doc.add_table(rows=5, cols=2)
        self._set_table_borders_none(sig)
        self._set_table_column_widths(sig, [3.5, 2.5])
        self._set_table_alignment(sig, "left")

        def _cell_lines(cell, lines):
            cell.paragraphs[0].clear()
            first = True
            for text in lines:
                paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                run = paragraph.add_run(text)
                self._apply_run_style(run, size=12, bold=False)

        _cell_lines(sig.rows[0].cells[0], ["___________________________"])
        _cell_lines(sig.rows[0].cells[1], ["___________________________"])
        _cell_lines(sig.rows[1].cells[0], [f"Name: {internal_name}".rstrip()])
        _cell_lines(sig.rows[1].cells[1], [f"Name: {external_name}".rstrip()])
        _cell_lines(sig.rows[2].cells[0], [f"Designation: {internal_desig}".rstrip()])
        _cell_lines(sig.rows[2].cells[1], [f"Designation: {external_desig}".rstrip()])
        _cell_lines(sig.rows[3].cells[0], [""])
        _cell_lines(sig.rows[3].cells[1], [f"Organization: {external_org}".rstrip()])
        _cell_lines(sig.rows[4].cells[0], ["(Internal Advisor)"])
        _cell_lines(sig.rows[4].cells[1], ["(External Advisor)"])

        hod_spacer = self.doc.add_paragraph()
        hod_spacer.paragraph_format.space_before = Pt(36)
        hod_line = self.doc.add_paragraph()
        hod_line.paragraph_format.space_after = Pt(2)
        hod_run = hod_line.add_run("___________________________")
        self._apply_run_style(hod_run, size=12, bold=False)
        if hod_name:
            hod_name_para = self.doc.add_paragraph()
            hod_name_para.paragraph_format.space_after = Pt(2)
            hr = hod_name_para.add_run(f"Name: {hod_name}")
            self._apply_run_style(hr, size=12, bold=False)
        hod_label = self.doc.add_paragraph()
        hod_label.paragraph_format.space_after = Pt(0)
        hl = hod_label.add_run("(Head of the Department)")
        self._apply_run_style(hl, size=12, bold=False)

        self.doc.add_page_break()
        self.current_chapter = previous_chapter

    def _normalize_students(self, meta):
        raw = meta.get("students") or meta.get("teamMembers") or []
        students = []
        if isinstance(raw, list):
            for item in raw:
                if isinstance(item, dict):
                    students.append(
                        {
                            "name": str(item.get("name") or "").strip(),
                            "enrollment": str(
                                item.get("enrollment") or item.get("enrolment") or ""
                            ).strip(),
                            "seatNumber": str(
                                item.get("seatNumber") or item.get("seat") or ""
                            ).strip(),
                        }
                    )
                elif str(item).strip():
                    students.append(
                        {"name": str(item).strip(), "enrollment": "", "seatNumber": ""}
                    )
        while len(students) < 4:
            students.append({"name": "", "enrollment": "", "seatNumber": ""})
        return students[:4]

    def _split_title_lines(self, title, max_lines=2):
        text = " ".join(str(title).split())
        if not text:
            return ["PROJECT TITLE"]
        if max_lines < 2 or len(text) <= 42:
            return [text]
        words = text.split()
        mid = (len(words) + 1) // 2
        return [" ".join(words[:mid]), " ".join(words[mid:])]

    def _resolve_logo_path(self, meta):
        """Resolve uploaded university logo from path or data URL."""
        candidates = []
        logo_path = meta.get("logoPath")
        if logo_path:
            candidates.append(Path(str(logo_path)))
        data_url = meta.get("logoDataUrl")
        uni_logo = meta.get("universityLogo")
        if not data_url and isinstance(uni_logo, dict):
            data_url = uni_logo.get("dataUrl")
            # Also accept a persisted file path if present.
            if uni_logo.get("path"):
                candidates.append(Path(str(uni_logo["path"])))
        for path in candidates:
            if path.exists() and path.stat().st_size > 0:
                return path
        if isinstance(data_url, str) and data_url.startswith("data:") and ";base64," in data_url:
            header, b64 = data_url.split(";base64,", 1)
            ext = ".png"
            lower = header.lower()
            if "jpeg" in lower or "jpg" in lower:
                ext = ".jpg"
            elif "gif" in lower:
                ext = ".gif"
            elif "webp" in lower:
                ext = ".webp"
            import base64
            import tempfile

            raw = base64.b64decode(b64)
            if not raw:
                return None
            handle = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
            handle.write(raw)
            handle.close()
            path = Path(handle.name)
            if path.exists() and path.stat().st_size > 0:
                return path
        return None

    def _add_logo_or_reserved_space(
        self, meta, width_inches, height_inches, space_before_pt=0, space_after_pt=0
    ):
        logo_path = self._resolve_logo_path(meta)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(space_before_pt)
        paragraph.paragraph_format.space_after = Pt(space_after_pt)
        paragraph.paragraph_format.line_spacing = 1.0
        if logo_path is not None:
            run = paragraph.add_run()
            run.add_picture(
                str(logo_path),
                width=Inches(width_inches),
                height=Inches(height_inches),
            )
        else:
            # Keep vertical space so layout does not collapse without a logo.
            paragraph.paragraph_format.space_before = Pt(
                space_before_pt + (height_inches * 72) / 2
            )
            paragraph.paragraph_format.space_after = Pt(
                space_after_pt + (height_inches * 72) / 2
            )
        return paragraph

    def _apply_run_style(self, run, size=12, bold=False, small_caps=False, underline=False):
        self.force_times_new_roman(run)
        run.font.size = Pt(size)
        run.bold = bold
        run.font.small_caps = small_caps
        run.underline = underline

    def _add_centered_line(
        self,
        text,
        size=12,
        bold=False,
        space_after=6,
        line_spacing=1.0,
        small_caps=False,
    ):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(space_after)
        paragraph.paragraph_format.line_spacing = line_spacing
        run = paragraph.add_run(str(text))
        self._apply_run_style(run, size=size, bold=bold, small_caps=small_caps)
        return paragraph

    def _set_paragraph_tab_stops(self, paragraph, positions_inches):
        p_pr = paragraph._element.get_or_add_pPr()
        for child in list(p_pr):
            if child.tag == qn("w:tabs"):
                p_pr.remove(child)
        tabs = OxmlElement("w:tabs")
        for pos in positions_inches:
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "left")
            tab.set(qn("w:pos"), str(int(round(float(pos) * 1440))))
            tabs.append(tab)
        p_pr.append(tabs)

    def _add_student_tab_rows(
        self,
        students,
        tab_positions_inches,
        font_size=13,
        header_bold=True,
        placeholder=True,
        seat_header="Seat number",
    ):
        """Official Name / Enrolment / Seat columns via fixed tab stops (not spaces)."""
        headers = ["Name", "Enrolment", seat_header]
        any_real = any(
            s.get("name") or s.get("enrollment") or s.get("seatNumber") for s in students
        )

        def add_row(values, bold):
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0
            # Keep official tab columns at the left margin (no body first-line indent).
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            self._set_paragraph_tab_stops(paragraph, tab_positions_inches)
            for index, value in enumerate(values):
                if index:
                    paragraph.add_run("\t")
                run = paragraph.add_run(str(value))
                self._apply_run_style(run, size=font_size, bold=bold)

        add_row(headers, bold=header_bold)
        for student in students:
            has_any = bool(
                student.get("name") or student.get("enrollment") or student.get("seatNumber")
            )
            if has_any:
                values = [
                    student.get("name") or "",
                    student.get("enrollment") or "",
                    student.get("seatNumber") or "",
                ]
                add_row(values, bold=False)
            elif placeholder and not any_real:
                add_row(
                    [
                        "Name",
                        "Enrolment",
                        "Seat number" if seat_header == "Seat number" else "Seat",
                    ],
                    bold=header_bold,
                )
            else:
                add_row(["", "", ""], bold=False)

    def _set_table_borders_none(self, table):
        tbl = table._tbl
        tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        if tbl.tblPr is None:
            tbl.insert(0, tbl_pr)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            element.set(qn("w:sz"), "0")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "auto")
            borders.append(element)
        existing = tbl_pr.find(qn("w:tblBorders"))
        if existing is not None:
            tbl_pr.remove(existing)
        tbl_pr.append(borders)

    def _set_table_column_widths(self, table, widths_inches):
        table.autofit = False
        tbl = table._tbl
        tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        if tbl.tblPr is None:
            tbl.insert(0, tbl_pr)
        existing = tbl_pr.find(qn("w:tblW"))
        if existing is not None:
            tbl_pr.remove(existing)
        total = OxmlElement("w:tblW")
        total.set(qn("w:w"), str(int(round(sum(widths_inches) * 1440))))
        total.set(qn("w:type"), "dxa")
        tbl_pr.append(total)
        grid = tbl.find(qn("w:tblGrid"))
        if grid is not None:
            tbl.remove(grid)
        grid = OxmlElement("w:tblGrid")
        for width in widths_inches:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(int(round(width * 1440))))
            grid.append(col)
        tbl.insert(1, grid)
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                if index >= len(widths_inches):
                    continue
                tc_pr = cell._tc.get_or_add_tcPr()
                for child in list(tc_pr):
                    if child.tag == qn("w:tcW"):
                        tc_pr.remove(child)
                tc_w = OxmlElement("w:tcW")
                tc_w.set(qn("w:w"), str(int(round(widths_inches[index] * 1440))))
                tc_w.set(qn("w:type"), "dxa")
                tc_pr.append(tc_w)

    def _set_table_alignment(self, table, alignment="left"):
        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER if alignment == "center" else WD_TABLE_ALIGNMENT.LEFT
        )
        tbl_pr = table._tbl.tblPr
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            table._tbl.insert(0, tbl_pr)
        jc = tbl_pr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            tbl_pr.append(jc)
        jc.set(qn("w:val"), "center" if alignment == "center" else "left")

    def _set_table_cell_margins(self, table, inches=0.0):
        tbl = table._tbl
        tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        if tbl.tblPr is None:
            tbl.insert(0, tbl_pr)
        existing = tbl_pr.find(qn("w:tblCellMar"))
        if existing is not None:
            tbl_pr.remove(existing)
        margins = OxmlElement("w:tblCellMar")
        twips = str(int(round(float(inches) * 1440)))
        for edge in ("top", "left", "bottom", "right"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:w"), twips)
            node.set(qn("w:type"), "dxa")
            margins.append(node)
        tbl_pr.append(margins)

    def _set_table_left_indent(self, table, inches):
        tbl = table._tbl
        tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        if tbl.tblPr is None:
            tbl.insert(0, tbl_pr)
        existing = tbl_pr.find(qn("w:tblInd"))
        if existing is not None:
            tbl_pr.remove(existing)
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), str(int(round(inches * 1440))))
        indent.set(qn("w:type"), "dxa")
        tbl_pr.append(indent)

    def add_table_of_contents(self):
        # Abstract already ends with a page break; do not add a trailing page-break
        # paragraph after the TOC field (that empty+break combo creates a blank page
        # when the TOC spans more than one page).
        self._add_generated_list_page(
            "TABLE OF CONTENTS",
            'TOC \\o "1-5" \\h \\z \\u',
            "Table of Contents will appear here after the document is opened in Word.",
            heading_style="JUW Front Matter",
            page_break_before=False,
            trailing_page_break=False,
        )

    def add_list_of_figures(self):
        self._add_generated_list_page(
            "LIST OF FIGURES",
            'TOC \\h \\z \\c "Figure"',
            "List of Figures will appear here after the document is opened in Word.",
            heading_style="JUW Front Matter",
            page_break_before=True,
            trailing_page_break=False,
        )

    def add_list_of_tables(self):
        self._add_generated_list_page(
            "LIST OF TABLES",
            'TOC \\h \\z \\c "Table"',
            "List of Tables will appear here after the document is opened in Word.",
            heading_style="JUW Front Matter",
            page_break_before=True,
            trailing_page_break=False,
        )

    def start_chapter(self, number, title):
        self.current_chapter = number
        self._start_body_section()
        self._clear_page_number_type()

        if number == 1:
            self._set_decimal_page_start(1)

        self.add_header_footer(f"Chapter {number}. {title}")
        self._increment_native_chapter()
        self.add_heading(f"CHAPTER {number}", 1)
        self.add_heading(title.upper(), 1)

    def start_unnumbered_section(self, title):
        self.current_chapter = 0
        self._start_body_section()
        self._clear_page_number_type()
        self.add_header_footer(title)
        self.add_heading(title.upper(), 1)

    def add_header_footer(self, chapter_title):
        section = self.doc.sections[-1]
        header = section.header
        header.is_linked_to_previous = False
        section.different_first_page_header_footer = False

        header_para = header.paragraphs[0]
        header_para.clear()
        header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_run = header_para.add_run(chapter_title)
        self.force_times_new_roman(header_run)
        header_run.font.size = Pt(10)
        header_run.italic = True

        footer_para = section.footer.paragraphs[0]
        footer_para.clear()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_run = insert_page(footer_para)
        self.force_times_new_roman(page_run)
        page_run.font.size = Pt(10)

    def add_figure(self, image_path, caption, width_percent=None, alignment=None):
        fig_cfg = self.config.get("figures", {})
        max_width = float(fig_cfg.get("width_inches", 5.5))
        try:
            pct = float(width_percent) if width_percent is not None else 100.0
        except (TypeError, ValueError):
            pct = 100.0
        pct = max(40.0, min(100.0, pct))
        width = max_width * (pct / 100.0)

        self.doc.add_picture(image_path, width=Inches(width))

        img_para = self.doc.paragraphs[-1]
        align = str(alignment or fig_cfg.get("alignment") or "center").lower()
        if align == "left":
            img_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif align == "right":
            img_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_para.paragraph_format.space_before = Pt(fig_cfg.get("space_before_image", 8))
        img_para.paragraph_format.space_after = Pt(fig_cfg.get("space_after_image", 2))

        paragraph = self.doc.add_paragraph()
        formatters.apply_caption_style(self.doc, paragraph)
        # Caption stays centered per JUW (numbering locked); image align is independent.
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_size = Pt(fig_cfg.get("caption_font_size", 12))

        prefix = paragraph.add_run(f"Figure {self.current_chapter}.")
        self.force_times_new_roman(prefix)
        prefix.font.size = font_size

        seq_run = insert_seq(paragraph, "Figure \\s 1")
        self.force_times_new_roman(seq_run)
        seq_run.font.size = font_size

        dot = paragraph.add_run(". ")
        self.force_times_new_roman(dot)
        dot.font.size = font_size

        caption_run = paragraph.add_run(caption)
        self.force_times_new_roman(caption_run)
        caption_run.font.size = font_size
        caption_run.bold = False
        caption_run.italic = False

        paragraph.paragraph_format.space_before = Pt(fig_cfg.get("space_before_caption", 3))
        paragraph.paragraph_format.space_after = Pt(fig_cfg.get("space_after_caption", 10))

    def add_table_caption(self, caption_text, chapter_num=None):
        chapter_num = self.current_chapter if chapter_num is None else chapter_num
        paragraph = self.doc.add_paragraph()
        formatters.apply_caption_style(self.doc, paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        prefix = paragraph.add_run(f"Table {chapter_num}.")
        self.force_times_new_roman(prefix)
        prefix.italic = True

        seq_run = insert_seq(paragraph, "Table \\s 1")
        self.force_times_new_roman(seq_run)
        seq_run.italic = True

        text_run = paragraph.add_run(f". {caption_text}")
        self.force_times_new_roman(text_run)
        text_run.italic = True

    def add_custom_table(self, data, columns, caption=None, with_caption=False):
        if len(self.doc.paragraphs) > 0:
            self.doc.add_paragraph()

        # Normalize legacy string cells to dicts.
        def as_cell(value):
            if isinstance(value, dict):
                return value
            return {"text": str(value or "")}

        header = [as_cell(cell) for cell in (columns or [])]
        body = [[as_cell(cell) for cell in row] for row in (data or [])]
        if not header:
            return None

        # Pad short body rows so merges stay within bounds.
        col_count = len(header)
        normalized_body = []
        for row in body:
            padded = list(row[:col_count])
            while len(padded) < col_count:
                padded.append({"text": ""})
            normalized_body.append(padded)
        body = normalized_body

        table = self.doc.add_table(rows=1 + len(body), cols=col_count)
        formatters.center_table(table)
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            tbl.append(tbl_pr)

        borders = OxmlElement("w:tblBorders")
        for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "000000")
            borders.append(border)
        tbl_pr.append(borders)

        grid = [header] + body
        for row_index, row_data in enumerate(grid):
            for col_index, cell_data in enumerate(row_data):
                if cell_data.get("hidden"):
                    continue
                cell = table.cell(row_index, col_index)
                self._apply_table_cell_content(
                    cell, cell_data, header=row_index == 0
                )

        # Apply merges after content so Word keeps the top-left cell text.
        for row_index, row_data in enumerate(grid):
            for col_index, cell_data in enumerate(row_data):
                if cell_data.get("hidden"):
                    continue
                colspan = int(cell_data.get("colspan") or 1)
                rowspan = int(cell_data.get("rowspan") or 1)
                if colspan <= 1 and rowspan <= 1:
                    continue
                end_row = min(row_index + rowspan - 1, len(grid) - 1)
                end_col = min(col_index + colspan - 1, len(header) - 1)
                try:
                    table.cell(row_index, col_index).merge(
                        table.cell(end_row, end_col)
                    )
                except Exception:
                    pass

        if with_caption or caption:
            self.add_table_caption(caption or "")
        return table

    def _apply_table_cell_content(self, cell, cell_data, header=False):
        text = str(cell_data.get("text") or "")
        cell.text = text
        paragraph = cell.paragraphs[0]
        run = paragraph.runs[0] if paragraph.runs else paragraph.add_run(text)
        if not paragraph.runs:
            run = paragraph.add_run(text)
        self.force_times_new_roman(run)

        bold = cell_data.get("bold")
        if bold is None:
            run.bold = True if header else False
        else:
            run.bold = bool(bold)

        align = cell_data.get("align")
        if align == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "left":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
            )

        color = cell_data.get("textColor") or cell_data.get("text_color")
        if color:
            hex_color = str(color).lstrip("#")
            if len(hex_color) == 6:
                try:
                    run.font.color.rgb = RGBColor.from_string(hex_color)
                except Exception:
                    pass

        background = cell_data.get("background")
        if background:
            hex_fill = str(background).lstrip("#")
            if len(hex_fill) == 6:
                tc_pr = cell._tc.get_or_add_tcPr()
                existing = tc_pr.find(qn("w:shd"))
                if existing is not None:
                    tc_pr.remove(existing)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), hex_fill.upper())
                tc_pr.append(shd)

    def _setup_title_cover_section(self):
        """Cover/title page: A4 margins, no page number."""
        section = self.doc.sections[0]
        formatters.apply_page_setup(section, self.config)
        section.different_first_page_header_footer = False
        footer = section.footer
        footer.is_linked_to_previous = False
        footer_para = footer.paragraphs[0]
        footer_para.clear()

    def _begin_roman_front_matter_section(self):
        """Start Approval..Ack section: Roman numerals; page i (Approval) has no number."""
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        formatters.apply_page_setup(section, self.config)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        formatters.restore_approval_page_without_number(self.doc, section=section)
        self._add_front_matter_footer(section)

    def _setup_front_matter_section(self):
        section = self.doc.sections[0]
        formatters.apply_page_setup(section, self.config)
        formatters.restore_approval_page_without_number(self.doc)
        self._add_front_matter_footer(section)

    def _add_front_matter_footer(self, section):
        footer = section.footer
        footer.is_linked_to_previous = False
        footer_para = footer.paragraphs[0]
        footer_para.clear()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_run = insert_page(footer_para)
        self.force_times_new_roman(page_run)
        page_run.font.size = Pt(self.config.get("styles", {}).get("page_number", {}).get("size", 10))

    def _add_generated_list_page(
        self,
        title,
        instruction,
        placeholder,
        heading_style="Heading 1",
        page_break_before=False,
        trailing_page_break=True,
    ):
        paragraph = self.doc.add_paragraph()
        self._apply_front_matter_heading(paragraph)
        if page_break_before:
            paragraph.paragraph_format.page_break_before = True
        self.add_bottom_border(paragraph)
        run = paragraph.add_run(title)
        self.force_times_new_roman(run)
        run.font.size = Pt(self._heading_size(1))
        run.bold = True

        field_para = self.doc.add_paragraph()
        field_para.style = self.doc.styles["Normal"]
        insert_field(field_para, instruction, placeholder)
        formatters.apply_body_paragraph(field_para, self.config, document=self.doc)
        if trailing_page_break:
            self.doc.add_page_break()

    def _start_body_section(self):
        new_section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        formatters.apply_page_setup(new_section, self.config)
        new_section.header.is_linked_to_previous = False
        new_section.footer.is_linked_to_previous = True
        return new_section

    def _clear_page_number_type(self):
        sect_pr = self.doc.sections[-1]._sectPr
        existing = sect_pr.find(qn("w:pgNumType"))
        if existing is not None:
            sect_pr.remove(existing)

    def _set_decimal_page_start(self, start):
        page_type = OxmlElement("w:pgNumType")
        page_type.set(qn("w:fmt"), "decimal")
        page_type.set(qn("w:start"), str(start))
        self.doc.sections[-1]._sectPr.append(page_type)

    def _increment_native_chapter(self):
        paragraph = self.doc.add_paragraph()
        p_pr = paragraph._element.get_or_add_pPr()

        num_pr = OxmlElement("w:numPr")
        ilvl_el = OxmlElement("w:ilvl")
        ilvl_el.set(qn("w:val"), "0")
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), "999")
        num_pr.append(ilvl_el)
        num_pr.append(num_id_el)
        p_pr.append(num_pr)

        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "2")
        spacing.set(qn("w:lineRule"), "exact")
        p_pr.append(spacing)

        r_pr = OxmlElement("w:rPr")
        r_pr.append(OxmlElement("w:vanish"))
        p_pr.append(r_pr)


def split_paragraphs(text):
    if not text:
        return []
    return [block.strip() for block in str(text).replace("\r\n", "\n").split("\n") if block.strip()]
