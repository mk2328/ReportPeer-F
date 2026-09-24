import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


def set_style_font_size(style, size_pt):
    """Write w:sz / w:szCs in half-points so python-docx and Word both see the size."""
    style.font.size = Pt(size_pt)
    r_pr = style.element.get_or_add_rPr()
    half_points = str(int(round(float(size_pt) * 2)))
    for tag in ("sz", "szCs"):
        element = r_pr.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            r_pr.append(element)
        element.set(qn("w:val"), half_points)


def set_style_toggle(r_pr, tag, enabled):
    element = r_pr.find(qn(f"w:{tag}"))
    if enabled:
        if element is None:
            element = OxmlElement(f"w:{tag}")
            r_pr.append(element)
        element.set(qn("w:val"), "true")
    elif element is not None:
        r_pr.remove(element)


def apply_page_setup(section, config):
    """Apply page size and JUW margins on a section."""
    layout = config.get("layout", {})
    page_size = str(layout.get("page_size", "A4")).strip().upper().replace(" ", "_")
    if page_size in {"LETTER", "US_LETTER", "US-LETTER"}:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    else:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    margins = layout.get("margins_inches", {})
    section.left_margin = Inches(margins.get("left", 1.5))
    section.right_margin = Inches(margins.get("right", 1.0))
    section.top_margin = Inches(margins.get("top", 1.0))
    section.bottom_margin = Inches(margins.get("bottom", 1.0))


def configure_body_style(document, config):
    """Normal style: Times New Roman 12pt, justified, 1.5 line spacing."""
    body = config.get("styles", {}).get("body", {})
    spacing = config.get("layout", {}).get("spacing", {})
    style = document.styles["Normal"]
    font_name = body.get("font", "Times New Roman")
    size_pt = body.get("size", 12)

    style.font.name = font_name
    style.font.bold = False
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    set_style_font_size(style, size_pt)

    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), font_name)

    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = spacing.get("line_spacing", 1.5)
    paragraph_format.space_before = Pt(spacing.get("paragraph_before", 0))
    paragraph_format.space_after = Pt(spacing.get("paragraph_after", 0))
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_pr = style.element.get_or_add_pPr()
    jc = p_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        p_pr.append(jc)
    jc.set(qn("w:val"), "both")


def setup_true_multilevel_list(document):
    """Inject native Word multilevel list numId 999 (chapter + 1.1 / 1.1.1)."""
    try:
        num_part = document.part.numbering_part
        if num_part is None:
            return

        abstract_num_xml = """
        <w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="999">
            <w:multiLevelType w:val="multilevel"/>
            <w:lvl w:ilvl="0">
                <w:start w:val="1"/>
                <w:numFmt w:val="decimal"/>
                <w:lvlText w:val="%1."/>
                <w:lvlJc w:val="left"/>
                <w:pPr><w:ind w:left="0" w:hanging="0"/></w:pPr>
            </w:lvl>
            <w:lvl w:ilvl="1">
                <w:start w:val="1"/>
                <w:numFmt w:val="decimal"/>
                <w:lvlText w:val="%1.%2"/>
                <w:lvlJc w:val="left"/>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                    <w:sz w:val="32"/>
                    <w:b/>
                </w:rPr>
                <w:pPr><w:ind w:left="720" w:hanging="720"/></w:pPr>
                <w:suff w:val="tab"/>
            </w:lvl>
            <w:lvl w:ilvl="2">
                <w:start w:val="1"/>
                <w:numFmt w:val="decimal"/>
                <w:lvlText w:val="%1.%2.%3"/>
                <w:lvlJc w:val="left"/>
                <w:pPr><w:ind w:left="1080" w:hanging="1080"/></w:pPr>
                <w:suff w:val="tab"/>
            </w:lvl>
            <w:lvl w:ilvl="3">
                <w:start w:val="1"/>
                <w:numFmt w:val="decimal"/>
                <w:lvlText w:val="%1.%2.%3.%4"/>
                <w:lvlJc w:val="left"/>
                <w:pPr><w:ind w:left="1440" w:hanging="1440"/></w:pPr>
                <w:suff w:val="tab"/>
            </w:lvl>
            <w:lvl w:ilvl="4">
                <w:start w:val="1"/>
                <w:numFmt w:val="decimal"/>
                <w:lvlText w:val="%1.%2.%3.%4.%5"/>
                <w:lvlJc w:val="left"/>
                <w:pPr><w:ind w:left="1800" w:hanging="1800"/></w:pPr>
                <w:suff w:val="tab"/>
            </w:lvl>
        </w:abstractNum>
        """
        num_xml = """
        <w:num xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:numId="999">
            <w:abstractNumId w:val="999"/>
        </w:num>
        """
        num_part._element.append(parse_xml(abstract_num_xml))
        num_part._element.append(parse_xml(num_xml))
    except Exception:
        pass


def cleanup_template(document):
    while len(document.paragraphs) > 1 and document.paragraphs[0].text.strip() == "":
        paragraph = document.paragraphs[0]
        paragraph._element.getparent().remove(paragraph._element)

    if document.paragraphs:
        document.paragraphs[0].style = document.styles["Normal"]
        document.paragraphs[0].clear()

    for index in range(1, 6):
        try:
            style = document.styles[f"Heading {index}"]
            p_pr = style.element.get_or_add_pPr()
            num_pr = p_pr.find(qn("w:numPr"))
            if num_pr is not None:
                p_pr.remove(num_pr)
        except Exception:
            pass


FRONT_MATTER_H1_TITLES = {
    "PROJECT APPROVAL FORM",
    "ABSTRACT",
    "TABLE OF CONTENTS",
    "LIST OF FIGURES",
    "LIST OF TABLES",
    "ACKNOWLEDGEMENT",
    "ACKNOWLEDGEMENTS",
}


def apply_heading_run(run, size_pt, bold=True, all_caps=False, small_caps=False):
    run.font.bold = bold
    run.font.all_caps = all_caps
    run.font.small_caps = small_caps
    apply_body_run(run, size_pt)
    r_pr = run._element.get_or_add_rPr()
    set_style_toggle(r_pr, "b", bold)
    set_style_toggle(r_pr, "caps", all_caps)
    set_style_toggle(r_pr, "smallCaps", small_caps)


def juw_heading_appearance(text, level):
    """
    JUW visual sizes. Heading 1 stays Heading 1 for TOC:
    Chapter N = 20pt; chapter name = 18pt; 1.1 / 1.1.1 / 1.1.1.1 = 16 / 14 / 12.
    """
    title = (text or "").strip().upper()
    if int(level or 1) == 1:
        if title.startswith("CHAPTER"):
            return 20, True, False
        if title in FRONT_MATTER_H1_TITLES:
            return 20, True, False
        return 18, True, False
    if int(level) == 2:
        return 16, False, True
    if int(level) == 3:
        return 14, False, True
    return 12, False, True


def enforce_juw_heading_runs(document):
    for paragraph in document.paragraphs:
        level = heading_level_from_style(paragraph)
        if not level:
            continue
        text = paragraph.text.strip()
        if not text:
            continue
        size_pt, all_caps, small_caps = juw_heading_appearance(text, level)
        if not paragraph.runs:
            continue
        for run in paragraph.runs:
            if not (run.text or "").strip():
                continue
            apply_heading_run(run, size_pt, bold=True, all_caps=all_caps, small_caps=small_caps)


def center_table(table):
    from docx.enum.table import WD_TABLE_ALIGNMENT

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    jc = tbl_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tbl_pr.append(jc)
    jc.set(qn("w:val"), "center")


def restore_approval_page_without_number(document, section=None):
    """Page i has no page number; Roman numbering still starts at 1 so Abstract is ii.

    When a title/cover section precedes approval, Roman settings belong on the
    front-matter section (usually sections[1]), not the unnumbered cover.
    """
    if section is None:
        section = document.sections[0]
        for candidate in document.sections:
            pg = candidate._sectPr.find(qn("w:pgNumType"))
            if pg is not None and pg.get(qn("w:fmt")) == "lowerRoman":
                section = candidate
                break
        else:
            # Title cover has no pgNumType; prefer the next section when present.
            if len(document.sections) > 1:
                first = document.sections[0]
                first_pg = first._sectPr.find(qn("w:pgNumType"))
                if first_pg is None and not first.footer.paragraphs[0].text.strip():
                    section = document.sections[1]

    section.different_first_page_header_footer = True
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    page_type = OxmlElement("w:pgNumType")
    page_type.set(qn("w:fmt"), "lowerRoman")
    page_type.set(qn("w:start"), "1")
    sect_pr.append(page_type)

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    for paragraph in first_footer.paragraphs:
        paragraph.clear()


def configure_front_matter_heading_style(document, config):
    """
    Front-matter titles look like Heading 1 but use outline level Body Text (9)
    so Word TOC \\o \"1-5\" does not list Approval/Abstract/TOC/LOF/LOT/Ack.
    """
    from docx.enum.style import WD_STYLE_TYPE

    name = "JUW Front Matter"
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]

    size_pt = config.get("styles", {}).get("headings", {}).get("JUW_H1", {}).get("size", 20)
    font = style.font
    font.name = "Times New Roman"
    font.bold = True
    font.color.rgb = RGBColor(0, 0, 0)
    font.all_caps = True
    font.italic = False
    set_style_font_size(style, size_pt)

    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "Times New Roman")
    r_fonts.set(qn("w:cs"), "Times New Roman")
    set_style_toggle(r_pr, "b", True)
    set_style_toggle(r_pr, "i", False)
    set_style_toggle(r_pr, "caps", True)

    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.5
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_pr = style.element.get_or_add_pPr()
    existing = p_pr.find(qn("w:outlineLvl"))
    if existing is not None:
        p_pr.remove(existing)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    p_pr.append(outline)
    return style


def configure_heading_styles(document, config):
    headings = config.get("styles", {}).get("headings", {})
    for level in range(1, 6):
        cfg = headings.get(f"JUW_H{level}", {})
        style = document.styles[f"Heading {level}"]
        size_pt = cfg.get("size", 20 - ((level - 1) * 2))
        font = style.font
        font.name = "Times New Roman"
        font.bold = cfg.get("bold", True)
        font.color.rgb = RGBColor(0, 0, 0)
        font.all_caps = cfg.get("all_caps", False)
        font.small_caps = cfg.get("small_caps", False)
        font.italic = False
        set_style_font_size(style, size_pt)

        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        r_fonts.set(qn("w:ascii"), "Times New Roman")
        r_fonts.set(qn("w:hAnsi"), "Times New Roman")
        r_fonts.set(qn("w:eastAsia"), "Times New Roman")
        r_fonts.set(qn("w:cs"), "Times New Roman")
        set_style_toggle(r_pr, "b", cfg.get("bold", True))
        set_style_toggle(r_pr, "i", False)
        set_style_toggle(r_pr, "caps", cfg.get("all_caps", False))
        set_style_toggle(r_pr, "smallCaps", cfg.get("small_caps", False))

        paragraph_format = style.paragraph_format
        paragraph_format.line_spacing = 1.5
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def configure_toc_styles(document):
    """TOC entries: Times New Roman 10 pt, single line spacing."""
    from docx.enum.text import WD_LINE_SPACING

    for index in range(1, 6):
        style = None
        for candidate in (f"TOC {index}", f"toc {index}"):
            try:
                style = document.styles[candidate]
                break
            except KeyError:
                continue
        if style is None:
            continue

        style.font.name = "Times New Roman"
        style.font.size = Pt(10)
        set_style_font_size(style, 10)

        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        r_fonts.set(qn("w:ascii"), "Times New Roman")
        r_fonts.set(qn("w:hAnsi"), "Times New Roman")
        r_fonts.set(qn("w:cs"), "Times New Roman")
        r_fonts.set(qn("w:eastAsia"), "Times New Roman")

        paragraph_format = style.paragraph_format
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph_format.line_spacing = 1.0
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)


def enforce_toc_entry_formatting(document):
    """Force generated TOC paragraphs to 10 pt TNR single spacing after Word updates."""
    from docx.enum.text import WD_LINE_SPACING

    for paragraph in document.paragraphs:
        name = (paragraph.style.name or "") if paragraph.style else ""
        if not name.lower().startswith("toc"):
            continue
        if not paragraph.text.strip():
            continue

        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        # TOC hyperlink fields often nest runs; size every w:r under the paragraph.
        for run_el in paragraph._p.iter(qn("w:r")):
            r_pr = run_el.find(qn("w:rPr"))
            if r_pr is None:
                r_pr = OxmlElement("w:rPr")
                run_el.insert(0, r_pr)
            r_fonts = r_pr.find(qn("w:rFonts"))
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                r_fonts.set(qn(f"w:{key}"), "Times New Roman")
            for tag in ("sz", "szCs"):
                element = r_pr.find(qn(f"w:{tag}"))
                if element is None:
                    element = OxmlElement(f"w:{tag}")
                    r_pr.append(element)
                element.set(qn("w:val"), "20")  # 10 pt in half-points



def apply_caption_style(document, paragraph):
    try:
        style = document.styles["Caption"]
    except KeyError:
        style = document.styles.add_style("Caption", 1)

    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.italic = True
    paragraph.style = style


# Foundational JUW body indent: 0.5" under H1/H2, 0.75" under H3+.
JUW_BODY_FIRST_LINE_H12_IN = 0.5
JUW_BODY_FIRST_LINE_H3_IN = 0.75


def is_body_paragraph(paragraph):
    if not paragraph.style:
        return False
    name = paragraph.style.name or ""
    lowered = name.lower()
    if paragraph_is_heading_name(name):
        return False
    if lowered.startswith("toc") or lowered.startswith("table of") or "caption" in lowered:
        return False
    return name == "Normal"


def paragraph_is_heading_name(name):
    return bool(name) and name.lower().startswith("heading")


def heading_level_from_style(paragraph):
    name = paragraph.style.name if paragraph.style else ""
    if not paragraph_is_heading_name(name):
        return None
    try:
        return int(name.split()[-1])
    except (TypeError, ValueError):
        return None


def is_list_paragraph(paragraph):
    text = (paragraph.text or "").lstrip()
    if text.startswith("\u2022") or text.startswith("o\t") or text.startswith("\u25AA"):
        return True
    if re.match(r"^\d+(?:\.\d+)*\.?(?:\t| )", text):
        return True
    first = paragraph.paragraph_format.first_line_indent
    return first is not None and first < 0


def apply_juw_body_indent(paragraph, heading_level=2):
    """First-line indent only, matching foundational add_body_text()."""
    inches = JUW_BODY_FIRST_LINE_H3_IN if int(heading_level or 2) >= 3 else JUW_BODY_FIRST_LINE_H12_IN
    paragraph.paragraph_format.left_indent = Inches(0)
    paragraph.paragraph_format.first_line_indent = Inches(inches)


def apply_body_run(run, size_pt=12):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), "Times New Roman")
    half_points = str(int(round(float(size_pt) * 2)))
    for tag in ("sz", "szCs"):
        element = r_pr.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            r_pr.append(element)
        element.set(qn("w:val"), half_points)


def apply_body_paragraph(paragraph, config, document=None):
    body = config.get("styles", {}).get("body", {})
    spacing = config.get("layout", {}).get("spacing", {})
    size_pt = body.get("size", 12)
    if document is not None:
        paragraph.style = document.styles["Normal"]
    paragraph.paragraph_format.line_spacing = spacing.get("line_spacing", 1.5)
    paragraph.paragraph_format.space_before = Pt(spacing.get("paragraph_before", 0))
    paragraph.paragraph_format.space_after = Pt(spacing.get("paragraph_after", 0))
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_pr = paragraph._element.get_or_add_pPr()
    jc = p_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        p_pr.append(jc)
    jc.set(qn("w:val"), "both")
    for run in paragraph.runs:
        apply_body_run(run, size_pt)


def _is_front_matter_form_paragraph(paragraph):
    """Cover/approval layout lines must not receive body indent or justify."""
    text = (paragraph.text or "").strip()
    if not text:
        return False
    # Official student rows use tab stops (Name / Enrolment / Seat).
    if "\t" in (paragraph.text or ""):
        return True
    if text.startswith("Signature:") or text.startswith("Name:") or text.startswith(
        "Designation:"
    ) or text.startswith("Organization:") or text.startswith("Project Title:"):
        return True
    if text.startswith("Project Advisor"):
        return True
    if text in {
        "By",
        "Submitted by",
        "Project Advisor",
        "Approval Committee:",
        "PROJECT APPROVAL",
        "Name",
        "Enrolment",
        "Seat",
        "Seat number",
    }:
        return True
    if text.endswith(":") and len(text) < 40:
        return True
    if text.startswith("(") and text.endswith(")"):
        return True
    if text.startswith("________________"):
        return True
    return False


def _table_has_left_indent(table):
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return False
    return tbl_pr.find(qn("w:tblInd")) is not None


def _is_front_matter_layout_table(table):
    """Approval header/signature tables must keep template column positions."""
    if _table_has_left_indent(table):
        return True
    texts = []
    for row in table.rows:
        for cell in row.cells:
            value = (cell.text or "").strip()
            if value:
                texts.append(value)
    joined = " ".join(texts)
    return (
        "Internal Advisor" in joined
        or "External Advisor" in joined
        or "DEPARTMENT OF COMPUTER SCIENCE" in joined.upper()
        or "JINNAH UNIVERSITY FOR WOMEN" in joined.upper()
    )


def enforce_body_paragraphs(document, config):
    """Force 12pt TNR, justify, and foundational first-line indent on body text."""
    heading_level = 2
    body = config.get("styles", {}).get("body", {})
    size_pt = body.get("size", 12)
    for paragraph in document.paragraphs:
        level = heading_level_from_style(paragraph)
        if level:
            heading_level = level
            continue
        if not paragraph.text.strip():
            continue
        if not is_body_paragraph(paragraph) or is_list_paragraph(paragraph):
            continue
        # Preserve title-page centering/sizes and approval signature/form lines.
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            continue
        if _is_front_matter_form_paragraph(paragraph):
            continue
        apply_body_paragraph(paragraph, config, document=document)
        apply_juw_body_indent(paragraph, heading_level)


def reapply_generated_styles(doc_path, config):
    """Re-apply JUW styles after Word field updates rewrite the file."""
    from docx import Document

    from core.file_io import save_docx_replacing, wait_until_unlocked

    wait_until_unlocked(str(doc_path))
    document = Document(str(doc_path))
    configure_body_style(document, config)
    configure_heading_styles(document, config)
    configure_front_matter_heading_style(document, config)
    configure_toc_styles(document)
    for section in document.sections:
        apply_page_setup(section, config)
    restore_approval_page_without_number(document)
    enforce_juw_heading_runs(document)
    enforce_body_paragraphs(document, config)
    enforce_toc_entry_formatting(document)
    for table in document.tables:
        # Do not center front-matter logo/signature tables — positions are template-fixed.
        if _is_front_matter_layout_table(table):
            continue
        center_table(table)
    save_docx_replacing(document, str(doc_path))
