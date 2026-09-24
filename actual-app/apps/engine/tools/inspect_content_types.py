"""Inspect a small generated DOCX for lists, tables, figures, and TOC/LOF/LOT."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

DOCX_PATH = Path(__file__).with_name("generated_content_types_test.docx")


def text_of(el) -> str:
    return "".join(node.text or "" for node in el.iter(qn("w:t")))


def has_field(paragraph, token: str) -> bool:
    xml = paragraph._p.xml
    return token in xml


def inspect(path: Path) -> dict:
    document = Document(str(path))
    paragraphs = list(document.paragraphs)
    texts = [p.text.strip() for p in paragraphs if p.text.strip()]

    bullet_paras = []
    number_paras = []
    markdown_like = []
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        num_pr = paragraph._p.find(qn("w:pPr"))
        has_num = False
        if num_pr is not None and num_pr.find(qn("w:numPr")) is not None:
            has_num = True
        if text[:1] in {"•", "o", "▪", "-", "*"} or text.startswith("- ") or text.startswith("* "):
            markdown_like.append(text[:80])
        if has_num:
            number_paras.append(text[:80])
        if text.startswith("•") or "\u2022" in text[:2]:
            bullet_paras.append(text[:80])

    toc = any("TABLE OF CONTENTS" in t.upper() for t in texts)
    lof = any("LIST OF FIGURES" in t.upper() for t in texts)
    lot = any("LIST OF TABLES" in t.upper() for t in texts)
    toc_field = any(has_field(p, "TOC ") for p in paragraphs)
    lof_field = any('TOC \\h \\z \\c "Figure"' in p._p.xml or 'c "Figure"' in p._p.xml for p in paragraphs)
    lot_field = any('c "Table"' in p._p.xml for p in paragraphs)

    figure_captions = [t for t in texts if t.upper().startswith("FIGURE ")]
    table_captions = [t for t in texts if t.upper().startswith("TABLE ") and "TABLE OF" not in t.upper()]

    caption_below_images = []
    body = document.element.body
    children = list(body.iterchildren())
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        if child.find(".//" + qn("w:drawing")) is None and child.find(".//" + qn("a:blip")) is None:
            has_blip = "a:blip" in child.xml or "w:drawing" in child.xml
            if not has_blip:
                continue
        next_text = ""
        for nxt in children[index + 1 :]:
            if nxt.tag == qn("w:p"):
                next_text = text_of(nxt).strip()
                if next_text:
                    break
        caption_below_images.append(next_text.lower().startswith("figure "))

    caption_below_tables = []
    for table in document.tables:
        nxt = table._tbl.getnext()
        nxt_text = text_of(nxt).strip() if nxt is not None else ""
        prev = table._tbl.getprevious()
        prev_text = text_of(prev).strip() if prev is not None else ""
        caption_below_tables.append(
            {
                "caption_after": nxt_text.lower().startswith("table "),
                "caption_before": prev_text.lower().startswith("table "),
                "after": nxt_text[:80],
                "before": prev_text[:80],
            }
        )

    images = len(document.inline_shapes)
    widths = []
    for shape in document.inline_shapes:
        try:
            widths.append(round(shape.width.inches, 2))
        except Exception:
            widths.append(None)

    bullets = [t for t in texts if t.startswith("\u2022") or t.startswith("•")]
    numbered = [t for t in texts if t[:2] in {"1.", "2."} and "markdown" not in t.lower() and "\t" in t or t.startswith("1.") or t.startswith("2.")]
    numbered = [t for t in texts if t.startswith("1.") or t.startswith("2.")]
    numbered = [t for t in numbered if "CHAPTER" not in t.upper() and not t.startswith("1.1") and not t.startswith("1.2") and not t.startswith("1.3")]

    verdict = {
        "Bullets": "PASS" if len(bullets) >= 2 else "FAIL",
        "Numbered Lists": "PASS" if any("First numbered" in t for t in numbered) and any("Second numbered" in t for t in numbered) else "FAIL",
        "Tables": "PASS" if len(document.tables) >= 1 and any(pos["caption_after"] for pos in caption_below_tables) else "FAIL",
        "Images": "PASS" if images >= 2 else "FAIL",
        "Figure Captions": "PASS" if caption_below_images and all(caption_below_images) else "FAIL",
        "Figure Numbering": "PASS" if any("Figure 1.1" in t for t in figure_captions) and any("Figure 1.2" in t for t in figure_captions) else "FAIL",
        "TOC": "PASS" if toc and toc_field else "FAIL",
        "LOF": "PASS" if lof and lof_field and any("Figure 1." in t for t in texts) else "FAIL",
        "LOT": "PASS" if lot and lot_field and any("Table 1.1" in t for t in texts) else "FAIL",
    }

    return {
        "docx": str(path),
        "verdict": verdict,
        "paragraph_count": len(paragraphs),
        "bullets_native_or_glyph": bullet_paras,
        "numbered_native": number_paras,
        "markdown_like_plain_text": markdown_like,
        "images": images,
        "image_widths_in": widths,
        "figure_captions": figure_captions,
        "caption_immediately_below_each_image": caption_below_images,
        "tables": len(document.tables),
        "table_captions": table_captions,
        "table_caption_positions": caption_below_tables,
        "toc_heading": toc,
        "lof_heading": lof,
        "lot_heading": lot,
        "toc_field": toc_field,
        "lof_field": lof_field,
        "lot_field": lot_field,
        "sample_texts": texts[:30],
    }


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DOCX_PATH
    if not path.exists():
        print(json.dumps({"ok": False, "error": f"DOCX not found: {path}"}))
        return 2
    print(json.dumps(inspect(path), indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
